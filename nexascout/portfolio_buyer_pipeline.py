# -*- coding: utf-8 -*-
r"""
NexaScout Universal Autonomous Buyer Acquisition Pipeline
Uçtan Uca Otonom Portföy Alıcı Bulma Motoru:
  1. Portföy Linkini Scrape Eder (Feature X: PSI REST + Playwright + CB BS4 + Slug Fallback)
  2. Fotoğrafları, Açıklamayı ve Konumu Multimodal Ajanlarla İnceler (Feature Y: PhotoVision + LegalZoning + InstitutionalMatching)
  3. 8 Kurumsal Alıcı Sektörünü ve Karar Verici Personalarını Eşleştirir (PropFit Skoru 0-100)
  4. Meslektaş Ağı, WhatsApp C-Level, Kurumsal E-Posta ve Vitrin Afiş Metinlerini Üretir
  5. Kurumsal Word (.docx) Yatırım Sunum Dosyasını Oluşturur
"""

import os
import re
import json
from datetime import datetime
from typing import Dict, Any, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
try:
    from nexascout.portfolio_scraper import scrape_any_listing
    from nexascout.property_analyzer import analyze_scraped_property
except ImportError:
    try:
        from modules.portfolio_scraper import scrape_any_listing
        from modules.property_analyzer import analyze_scraped_property
    except ImportError:
        from portfolio_scraper import scrape_any_listing
        from property_analyzer import analyze_scraped_property

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PORTFOLIOS_DIR = os.path.join(BASE_DIR, "data", "PORTFOYLER")


def _generate_portfolio_docx(target_dir: str, analysis: Dict[str, Any], broker: Dict[str, str], url: str) -> str:
    """Portföy için şık ve resmi bir Word (.docx) yatırım teklif mektubu oluşturur."""
    docx_path = os.path.join(target_dir, "PORTFOY_YATIRIM_VE_HEDEF_ALICI_DOSYASI.docx")
    
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Üst Bilgi
    h_p = doc.add_paragraph()
    h_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = h_p.add_run("MÜNHASIR YETKİLİ KURUMSAL YATIRIM TEKLİFİ\n")
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGBColor(197, 160, 89)  # Gold
    r2 = h_p.add_run(f"{broker.get('office', 'COLDWELL BANKER VIP REAL')}\nDanışman: {broker.get('name', 'Yiğit Narin')} ({broker.get('phone', '0532 451 40 08')})")
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(100, 116, 139)

    # Başlık
    t_p = doc.add_paragraph()
    t_run = t_p.add_run(f"\n{analysis['title'].upper()}\n")
    t_run.font.bold = True
    t_run.font.size = Pt(16)
    t_run.font.color.rgb = RGBColor(27, 54, 93)  # Navy Dark

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Konum: {analysis['location']}  |  Ölçek: {analysis['kapali_m2']} ({analysis['kat_sayisi']})  |  Fiyat: {analysis['price']}")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    # 1. Haksız Avantajlar
    doc.add_heading("1. Mülkün Stratejik Değerleri & Haksız Avantajları", level=2)
    for adv in analysis.get("advantages", []):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(adv)
        r.font.size = Pt(10.5)

    # 2. Görsel & Mimari Analiz Özeti
    vision = analysis.get("vision_analysis", {})
    if vision:
        doc.add_heading("2. Mimari Kondisyon & Görsel Değerlendirme", level=2)
        v_p = doc.add_paragraph()
        v_p.add_run(f"Genel Kondisyon: {vision.get('overall_condition', 'Bakımlı')} (Skor: {vision.get('condition_score', 9)}/10)\n").bold = True
        v_p.add_run(f"Ticari Uygunluk: {vision.get('commercial_suitability', 'Çok Yüksek')}\n")
        v_p.add_run(f"Doğal Aydınlatma & Cephe: {vision.get('natural_light', 'Ferah')}\n")
        v_p.add_run(f"Analiz Modeli: {vision.get('analyzed_with', 'NexaScout Vision Agent')}")

    # 3. Hedef Kurumsal Alıcı Sektörleri
    doc.add_heading("3. Öncelikli Kurumsal Alıcı Sektörleri & PropFit Skorları", level=2)
    
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Sektör & Alıcı Grubu"
    hdr_cells[1].text = "PropFit"
    hdr_cells[2].text = "Stratejik Uyum Gerekçesi"

    for c in hdr_cells:
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>')
        c._tc.get_or_add_tcPr().append(shd)

    for sec in analysis.get("matched_sectors", [])[:5]:
        row_cells = table.add_row().cells
        row_cells[0].text = sec["ad"]
        row_cells[1].text = f"%{sec['propfit_score']}"
        row_cells[2].text = sec["neden"]
        
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].paragraphs[0].runs[0].font.bold = True
        row_cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 0, 0)
        
        for c in row_cells:
            c.paragraphs[0].runs[0].font.size = Pt(9.5)

    # İletişim Bilgileri
    doc.add_paragraph("\n")
    c_p = doc.add_paragraph()
    c_p.add_run("Tek Yetkili Portföy Temsilcisi:\n").bold = True
    c_p.add_run(f"{broker.get('name', 'Yiğit Narin')} — {broker.get('office', 'Coldwell Banker VIP Real')}\n")
    c_p.add_run(f"Telefon: {broker.get('phone', '0532 451 40 08')}\n")
    c_p.add_run(f"İlan Detayı: {url}")

    doc.save(docx_path)
    return docx_path


def run_portfolio_buyer_pipeline(
    url: str,
    broker_name: str = "YİĞİT NARİN",
    broker_phone: str = "0532 451 40 08",
    broker_office: str = "Coldwell Banker VIP Real Gayrimenkul"
) -> Dict[str, Any]:
    """
    Tek bir link ile tüm alıcı bulma ve stratejik analiz sürecini otonom olarak yürütür.
    """
    # 1. Scrape (Feature X)
    scraped = scrape_any_listing(url)
    
    # 2. Çoklu Ajan Analizi (Feature Y: Vision + Legal + Institutional)
    analysis = analyze_scraped_property(scraped)

    # 3. Klasörleme
    listing_id = scraped.get("listing_id") or "PORTFOY"
    clean_title = re.sub(r"[^a-zA-Z0-9_\-]", "_", scraped.get("title", "Portfoy")[:35])
    folder_name = f"{listing_id}_{clean_title}".strip("_")
    
    target_dir = os.path.join(PORTFOLIOS_DIR, folder_name)
    os.makedirs(target_dir, exist_ok=True)

    broker_info = {"name": broker_name, "phone": broker_phone, "office": broker_office}

    # 4. Avantajlar Metni
    advantages_bullet = "\n".join(f"• {adv}" for adv in analysis["advantages"])

    # 5. Meslektaş Duyuru Metni
    meslektas_metni = f"""🚨 DEĞERLİ MESLEKTAŞLARIM
📍 {analysis['location'].upper()}
⭐ {analysis['title'].upper()}
{analysis['asset_class'].upper()} • {analysis['kapali_m2']}

🔑 MÜLKÜN EN BÜYÜK AVANTAJLARI (ALICI KOZLARI):
{advantages_bullet}

🎯 HANGİ MÜŞTERİLERİNİZE SUNABİLİRSİNİZ?:
{chr(10).join(f"✔ {s['ad']} (PropFit: %{s['propfit_score']}) — {s['neden']}" for s in analysis['matched_sectors'][:4])}

🤝 Meslektaş iş birliğine ve ortak çalışmaya tam açığız. Portföyünüzde bu profilde alıcı arayan müşteriniz varsa GETİRİN, BİRLİKTE GÖSTERELİM.

📞 {broker_name} — {broker_phone}
{broker_office}

Müşteriniz varsa, mülk hazır. 🔑
İlan: {url}
"""

    # 6. WhatsApp VIP Yönetici Teaser'ı
    wa_teaser = f"""Merhaba, iyi çalışmalar. Ben {broker_name}, {broker_office} Gayrimenkul ve Yatırım Danışmanıyım.

Kurumunuzun Ankara ve bölge büyüme hedefleri doğrultusunda değerlendirebileceğinizi düşündüğüm çok özel bir mülk için doğrudan şahsınıza ulaşmak istedim.

📍 {analysis['location']}
🏢 {analysis['asset_class']} | {analysis['kapali_m2']}
⭐ Öne Çıkan Değerler:
{chr(10).join(f"• {adv}" for adv in analysis['advantages'][:3])}

Kurumunuz için prestijli yeni hizmet noktası / karargâh olarak kullanıma hazır durumdadır.

İlgili yatırım veya gayrimenkul geliştirme biriminize yönlendirebilirseniz memnuniyet duyarım. Müsait olduğunuzda teknik şartname dosyasını paylaşmak ve yerinde göstermek isterim.

🔗 Resmi İlan Linki:
{url}

📞 {broker_name} — {broker_phone}
"""

    # 7. Kurumsal E-Posta Şablonu
    email_subject = f"Kurumsal Genişleme & Stratejik Yatırım Fırsatı | {analysis['location']}"
    email_body = f"""Sayın Yetkili,

Ben {broker_name}, {broker_office} Gayrimenkul ve Yatırım Danışmanıyım.

Kurumunuzun sektördeki büyüme vizyonunu ve hizmet ağını yakından takip etmekteyiz. Büyüme ve yeni yatırım planlarınız açısından yüksek katma değer sağlayacağını öngördüğümüz, bölgenin en stratejik noktasında yer alan özel bir portföyümüz için size doğrudan ulaşmaktayız.

MÜLK BİLGİLERİ & ÖNE ÇIKAN DEĞERLER:
• Lokasyon: {analysis['location']}
• Mülk Türü: {analysis['asset_class']} ({analysis['kapali_m2']})
{advantages_bullet}

ÖNCELİKLİ KULLANIM VE HİZMET ALANLARI:
{chr(10).join(f"• {s['ad']} (PropFit Uyumu: %{s['propfit_score']})" for s in analysis['matched_sectors'][:3])}

Portföyümüz tek yetki sözleşmesiyle temsil edilmekte olup, kurumunuzun ilgili gayrimenkul geliştirme / yatırım birimine iletilmesini rica ederiz. Uygun göreceğiniz bir zaman diliminde teknik sunum dosyasını aktarmaktan ve yerinde inceleme organize etmekten memnuniyet duyarız.

Resmi İlan Detayı:
{url}

Saygılarımla,

{broker_name}
{broker_office}
İletişim: {broker_phone}
"""

    # 8. Görsel / Afiş Vitrin Metni
    gorsel_metin = f"""🚨 {analysis['location'].upper()}
⭐ {analysis['title'].upper()}

✔ {analysis['asset_class']} ({analysis['kapali_m2']})
{chr(10).join(f"✔ {adv}" for adv in analysis['advantages'][:3])}

💼 Tek Yetkili Danışman: {broker_name}
📱 {broker_phone} | {broker_office}
🔗 {url}
"""

    # 9. Dosyaları Kaydet (.json, .txt, .docx, .md)
    with open(os.path.join(target_dir, "portfolio_data.json"), "w", encoding="utf-8") as f:
        json.dump({
            "scraped": scraped,
            "analysis": analysis,
            "broker": broker_info,
            "created_at": datetime.now().isoformat()
        }, f, ensure_ascii=False, indent=2)

    with open(os.path.join(target_dir, "MESLEKTAS_IS_BIRLIGI_DUYURUSU.txt"), "w", encoding="utf-8") as f:
        f.write(meslektas_metni)

    with open(os.path.join(target_dir, "VIP_WHATSAPP_TEASER.txt"), "w", encoding="utf-8") as f:
        f.write(wa_teaser)

    with open(os.path.join(target_dir, "KURUMSAL_EPOSTA.txt"), "w", encoding="utf-8") as f:
        f.write(f"KONU: {email_subject}\n\n{email_body}")

    with open(os.path.join(target_dir, "GORSEL_AFIS_METNI.txt"), "w", encoding="utf-8") as f:
        f.write(gorsel_metin)

    # Word (.docx) Dosyasını Oluştur
    docx_path = _generate_portfolio_docx(target_dir, analysis, broker_info, url)

    # 10. Stratejik Markdown Raporu
    rapor_md = f"""# 🏛️ STRATEJİK PORTFÖY ALICI BULMA VE YATIRIM RAPORU
> **Mülk:** {analysis['title']}  
> **Konum:** {analysis['location']}  
> **Sınıf:** {analysis['asset_class']} ({analysis['kapali_m2']})  
> **Tek Yetkili Danışman:** {broker_name} ({broker_phone}) — {broker_office}  
> **İlan Linki:** {url}  

---

## 📌 1. PORTFÖYÜN HAKSIZ AVANTAJLARI (UNFAIR ADVANTAGES)
{advantages_bullet}

---

## 🖼️ 2. MİMARİ VE GÖRSEL ANALİZ (VISION AGENT)
- **Genel Kondisyon:** {analysis['vision_analysis'].get('overall_condition', 'Bakımlı')}
- **Kondisyon Skoru:** {analysis['vision_analysis'].get('condition_score', 9)} / 10
- **Ticari Uygunluk:** {analysis['vision_analysis'].get('commercial_suitability', 'Çok Yüksek')}
- **Tespit Edilen Mimari Nitelikler:** {', '.join(analysis['vision_analysis'].get('detected_features', []))}

---

## 🎯 3. TESPİT EDİLEN ÖNCELİKLİ ALICI SEKTÖRLERİ & PROPFIT SKORLARI
| Sektör / Alıcı Profili | PropFit | Uyum Durumu | Eşleşme Gerekçesi |
| :--- | :---: | :---: | :--- |
{chr(10).join(f"| **{s['ad']}** | **%{s['propfit_score']}** | `{s['fit_verdict']}` | {s['neden']} |" for s in analysis['matched_sectors'])}

---

## 🔍 4. HARİTADAN HEDEFLENECEK KURUMSAL ARAMA SORGULARI
{chr(10).join(f"- `{q}`" for q in analysis['lead_search_queries'])}

---

## 💬 5. OLUŞTURULAN DOSYA VE MATERYALLER
Tüm çıktılar `{target_dir}` klasöründe oluşturulmuştur:
- 📄 [PORTFOY_YATIRIM_VE_HEDEF_ALICI_DOSYASI.docx](file:///{docx_path.replace(chr(92), '/')})
- 📝 [MESLEKTAS_IS_BIRLIGI_DUYURUSU.txt](file:///{target_dir.replace(chr(92), '/')}/MESLEKTAS_IS_BIRLIGI_DUYURUSU.txt)
- 📱 [VIP_WHATSAPP_TEASER.txt](file:///{target_dir.replace(chr(92), '/')}/VIP_WHATSAPP_TEASER.txt)
- ✉️ [KURUMSAL_EPOSTA.txt](file:///{target_dir.replace(chr(92), '/')}/KURUMSAL_EPOSTA.txt)
- 🎨 [GORSEL_AFIS_METNI.txt](file:///{target_dir.replace(chr(92), '/')}/GORSEL_AFIS_METNI.txt)
"""
    with open(os.path.join(target_dir, "STRATEJIK_RAPOR.md"), "w", encoding="utf-8") as f:
        f.write(rapor_md)

    return {
        "status": "success",
        "url": url,
        "scraped": scraped,
        "analysis": analysis,
        "folder_path": target_dir,
        "docx_path": docx_path,
        "outputs": {
            "meslektas_metni": meslektas_metni,
            "wa_teaser": wa_teaser,
            "email_subject": email_subject,
            "email_body": email_body,
            "gorsel_metin": gorsel_metin
        }
    }
